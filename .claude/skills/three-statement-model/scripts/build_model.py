#!/usr/bin/env python3
"""
Deterministic three-statement model builder.

Generates a fully linked, self-balancing 3-statement Excel model (income
statement, balance sheet, cash flow statement + working-capital, depreciation
and debt schedules, charts) from an inputs JSON, optionally adapted by a
jurisdiction pack (labels, units, default tax line, terminology notes).

The template it implements was reverse engineered from a proven industry
case-study workbook and is verified by scripts/validate_model.py (structure,
formula canon, values vs an independent simulation, recalculation).

Usage:
  python3 build_model.py --inputs inputs.json [--jurisdiction pack.json]
                         [--output model.xlsx]

Exit code 0 on success; non-zero with a message if input validation fails.
"""
import argparse
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef

BASE_LABELS = {
    5: "Assumptions", 7: "Key Assumptions",
    8: "Revenue Growth (% Change)", 9: "Cost of Goods Sold (% of Revenue)",
    10: "Salaries and Benefits (% of Revenue)", 11: "Rent and Overhead ({units})",
    12: "Depreciation & Amortization (% of PP&E)", 13: "Interest (% of Debt)",
    14: "Tax Rate (% of Earnings Before Tax)", 15: "Accounts Receivable (Days)",
    16: "Inventory (Days)", 17: "Accounts Payable (Days)",
    18: "Capital Expenditures ({units})", 19: "Debt Issuance (Repayment) ({units})",
    20: "Equity Issued (Repaid) ({units})", 21: "Days in Period",
    24: "Income Statement", 26: "Revenue", 27: "Cost of Goods Sold (COGS)",
    28: "Gross Profit", 29: "Expenses", 30: "Salaries and Benefits",
    31: "Rent and Overhead", 32: "Depreciation & Amortization", 33: "Interest",
    34: "Total Expenses", 35: "Earnings Before Tax", 37: "Taxes", 38: "Net Earnings",
    41: "Balance Sheet", 43: "Assets", 44: "Cash", 45: "Accounts Receivable",
    46: "Inventory", 47: "Property & Equipment", 48: "Total Assets",
    50: "Liabilities", 51: "Accounts Payable", 52: "Debt", 53: "Total Liabilities",
    54: "Shareholder's Equity", 55: "Equity Capital", 56: "Retained Earnings",
    57: "Shareholder's Equity", 58: "Total Liabilities & Shareholder's Equity",
    60: "Check",
    63: "Cash Flow Statement", 65: "Operating Cash Flow", 66: "Net Earnings",
    67: "Plus: Depreciation & Amortization", 68: "Less: Changes in Working Capital",
    69: "Cash from Operations", 71: "Investing Cash Flow",
    72: "Investments in Property & Equipment", 73: "Cash from Investing",
    75: "Financing Cash Flow", 76: "Issuance (repayment) of debt",
    77: "Issuance (repayment) of equity", 78: "Cash from Financing",
    80: "Net Increase (decrease) in Cash", 81: "Opening Cash Balance",
    82: "Closing Cash Balance",
    86: "Supporting Schedules", 88: "Working Capital Schedule",
    89: "Accounts Receivable", 90: "Inventory", 91: "Accounts Payable",
    92: "Net Working Capital (NWC)", 93: "Change in NWC",
    95: "Depreciation Schedule", 96: "PPE Opening", 97: "Plus Capex",
    98: "Less Depreciation", 99: "PPE Closing",
    101: "Debt & Interest Schedule", 102: "Debt Opening",
    103: "Issuance (repayment)", 104: "Debt Closing", 105: "Interest Expense",
    109: "Charts and Graphs", 111: "Revenue", 112: "Gross Profit Margin (%)",
    114: "Operating Cash Flow", 115: "Investing Cash Flow", 116: "Financing Cash Flow",
}
SECTION_ROWS = (5, 24, 41, 63, 86, 109)
SUBHEADER_ROWS = (7, 29, 43, 50, 54, 65, 71, 75, 88, 95, 101)
BOLD_TOTAL_ROWS = (28, 34, 35, 38, 48, 53, 57, 58, 69, 73, 78, 82)

NUM_FMT = '_-* #,##0_-;\\(#,##0\\)_-;_-* "-"_-;_-@'
PCT_FMT = '0.0%'
GPM_FMT = '0%'
CHK_FMT = '0.0000_ ;\\-0.0000\\ '


def resolve_labels(pack, payload=None):
    """Base labels + jurisdiction overrides. '{units}' expands to the payload's
    units_label if set, else the pack's, else $000's. Row 37 falls back to the
    pack's default_tax_line.label; explicit label_overrides win over both."""
    units = ((payload or {}).get("units_label")
             or (pack or {}).get("units_label", "$000's"))
    labels = {r: t.format(units=units) for r, t in BASE_LABELS.items()}
    tax_label = ((pack or {}).get("default_tax_line") or {}).get("label")
    if tax_label:
        labels[37] = tax_label
    for r, t in ((pack or {}).get("label_overrides") or {}).items():
        labels[int(r)] = t
    return labels


def apply_pack_defaults(payload, pack):
    """Fill missing forecast assumptions from the jurisdiction pack.
    Currently: tax_pct_ebt defaults to the pack's default tax line rate."""
    if not pack:
        return payload
    fa = payload["forecast_assumptions"]
    n = payload["n_forecast"]
    rate = (pack.get("default_tax_line") or {}).get("rate")
    if rate is not None and not fa.get("tax_pct_ebt"):
        fa["tax_pct_ebt"] = [rate] * n
    return payload


REQUIRED_TOP = ("company", "first_historical_year", "n_historical", "n_forecast",
                "historical", "forecast_assumptions")
REQUIRED_HIST = ("revenue", "cogs", "salaries_and_benefits", "rent_and_overhead",
                 "depreciation_amortization", "interest_expense", "taxes", "cash",
                 "accounts_receivable", "inventory", "ppe_closing",
                 "accounts_payable", "debt_closing", "equity_capital",
                 "retained_earnings", "net_earnings_cf", "depreciation_cf",
                 "change_in_nwc_cf", "capex", "debt_issuance_repayment",
                 "equity_issuance_repayment", "opening_cash_first_year",
                 "ppe_opening_first_year", "debt_opening_first_year")
REQUIRED_FA = ("revenue_growth_pct", "cogs_pct_revenue", "salaries_pct_revenue",
               "rent_and_overhead", "da_pct_opening_ppe", "interest_pct_avg_debt",
               "ar_days", "inventory_days", "ap_days", "capex",
               "debt_issuance_repayment", "equity_issuance_repayment")


def validate_inputs(payload):
    """§1 of the template spec: structure, lengths, four roll-forward
    identities (±1). Returns (errors, warnings)."""
    errors, warnings = [], []
    missing = [k for k in REQUIRED_TOP if k not in payload]
    if not missing:
        missing += [f"historical.{k}" for k in REQUIRED_HIST
                    if k not in payload["historical"]]
        missing += [f"forecast_assumptions.{k}" for k in REQUIRED_FA
                    if k not in payload["forecast_assumptions"]]
    if missing:
        return [f"missing required field: {k}" for k in missing], warnings
    H, F = payload["n_historical"], payload["n_forecast"]
    hist, fa = payload["historical"], payload["forecast_assumptions"]
    Y0 = payload["first_historical_year"]
    if H < 2:
        errors.append(f"n_historical must be >= 2, got {H}")
    if not fa.get("tax_pct_ebt"):
        errors.append("forecast_assumptions.tax_pct_ebt is missing/empty and the "
                      "jurisdiction pack provides no default rate - set it explicitly")
    if payload.get("days_in_period", 365) != 365:
        errors.append(f"days_in_period must be 365, got {payload.get('days_in_period')}")
    if H + F > 23:
        errors.append(f"horizon cap exceeded: {H}+{F} > 23")
    for k, v in hist.items():
        if isinstance(v, list) and len(v) != H:
            errors.append(f"historical.{k}: {len(v)} entries, expected {H}")
    for k, v in fa.items():
        if isinstance(v, list) and len(v) != F:
            errors.append(f"forecast_assumptions.{k}: {len(v)} entries, expected {F}")
    if errors:
        return errors, warnings
    cash_open = hist["opening_cash_first_year"]
    ppe_open = hist["ppe_opening_first_year"]
    debt_open = hist["debt_opening_first_year"]
    for i in range(H):
        yr = Y0 + i
        assets = (hist["cash"][i] + hist["accounts_receivable"][i]
                  + hist["inventory"][i] + hist["ppe_closing"][i])
        le = (hist["accounts_payable"][i] + hist["debt_closing"][i]
              + hist["equity_capital"][i] + hist["retained_earnings"][i])
        if abs(assets - le) > 1:
            errors.append(f"{yr}: BS identity broken: assets {assets:.2f} vs L+E {le:.2f}")
        cash_close = (cash_open + hist["net_earnings_cf"][i] + hist["depreciation_cf"][i]
                      - hist["change_in_nwc_cf"][i] - hist["capex"][i]
                      + hist["debt_issuance_repayment"][i]
                      + hist["equity_issuance_repayment"][i])
        if abs(cash_close - hist["cash"][i]) > 1:
            errors.append(f"{yr}: cash roll {cash_close:.2f} vs cash {hist['cash'][i]}")
        cash_open = hist["cash"][i]
        ppe_close = ppe_open + hist["capex"][i] - hist["depreciation_amortization"][i]
        if abs(ppe_close - hist["ppe_closing"][i]) > 1:
            errors.append(f"{yr}: PP&E roll {ppe_close:.2f} vs {hist['ppe_closing'][i]}")
        ppe_open = hist["ppe_closing"][i]
        debt_close = debt_open + hist["debt_issuance_repayment"][i]
        if abs(debt_close - hist["debt_closing"][i]) > 1:
            errors.append(f"{yr}: debt roll {debt_close:.2f} vs {hist['debt_closing'][i]}")
        debt_open = hist["debt_closing"][i]
        ne_is = (hist["revenue"][i] - hist["cogs"][i]
                 - hist["salaries_and_benefits"][i] - hist["rent_and_overhead"][i]
                 - hist["depreciation_amortization"][i]
                 - hist["interest_expense"][i] - hist["taxes"][i])
        if i > 0:
            if abs(hist["retained_earnings"][i - 1] + ne_is
                   - hist["retained_earnings"][i]) > 1:
                errors.append(f"{yr}: retained-earnings roll broken")
        # As-reported CFS lines should agree with their IS counterparts; a
        # large gap means the trial-balance mapping went wrong even though
        # the cash/RE identities can still both hold.
        if abs(hist["net_earnings_cf"][i] - ne_is) > 1:
            warnings.append(f"{yr}: CFS net earnings {hist['net_earnings_cf'][i]} "
                            f"differs from IS-derived {ne_is:.2f} - check mapping")
        if abs(hist["depreciation_cf"][i]
               - hist["depreciation_amortization"][i]) > 1:
            warnings.append(f"{yr}: CFS depreciation differs from IS D&A - "
                            f"check mapping")
    return errors, warnings


def populate_model_sheet(ws, payload, pack):
    """Populate a worksheet with the complete, validated single-sheet model
    (labels, blue inputs, black formulas, formatting, borders, conditional
    formats, the two charts). Factored out of build() so the identical canonical
    sheet can be embedded in a multi-sheet workbook (scripts/build_workbook.py)
    within a SINGLE save pass — openpyxl silently drops charts when it loads and
    re-saves a workbook, so every sheet must be written before the first save.
    Returns the column-map info dict. Does not touch workbook-level state."""
    H, F = payload["n_historical"], payload["n_forecast"]
    DIP = payload.get("days_in_period", 365)
    Y0 = payload["first_historical_year"]
    hist, fa = payload["historical"], payload["forecast_assumptions"]

    FIRST = 4
    LAST = FIRST + H + F - 1
    FC1 = FIRST + H
    HIST_COLS = list(range(FIRST, FC1))
    FC_COLS = list(range(FC1, LAST + 1))
    ALL_COLS = HIST_COLS + FC_COLS
    labels = resolve_labels(pack, payload)

    BODY = dict(name="Arial Narrow", size=12)
    F_BODY = Font(**BODY, color="FF000000")
    F_INPUT = Font(**BODY, color="FF0000FF")
    F_BODY_B = Font(**BODY, bold=True, color="FF000000")
    F_INPUT_B = Font(**BODY, bold=True, color="FF0000FF")
    F_SECTION = Font(name="Open Sans", size=14, bold=True, color="FF3271D2")
    F_YEAR_H = Font(name="Arial Narrow", size=14, bold=True, color="FF000000")
    F_YEAR_F = Font(name="Arial Narrow", size=14, bold=True, color="FFFFFFFF")
    F_CHECK = Font(name="Arial Narrow", size=10, italic=True, color="FF000000")
    FILL_LIGHT = PatternFill("solid", fgColor="FFE7F2FF")
    FILL_DARK = PatternFill("solid", fgColor="FF000C3F")

    for r, text in labels.items():
        cell = ws.cell(row=r, column=1, value=text)
        if r in SECTION_ROWS:
            cell.font = F_SECTION
            cell.fill = FILL_LIGHT
        elif r == 26 or r in SUBHEADER_ROWS or r in BOLD_TOTAL_ROWS:
            cell.font = F_BODY_B
        else:
            cell.font = F_BODY
    a3 = ws.cell(row=3, column=1, value="Balance Sheet Check")
    a3.font = F_CHECK

    ws.cell(row=2, column=FIRST, value=Y0)
    for c in ALL_COLS[1:]:
        ws.cell(row=2, column=c, value=f"=+{L(c-1)}2+1")
    ws.cell(row=1, column=FIRST, value="Historical Results")
    ws.cell(row=1, column=FC1, value=" Forecast Period")

    def put_input(row, col, value, bold=False):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = F_INPUT_B if bold else F_INPUT
        return cell

    def put_formula(row, col, formula, bold=False):
        cell = ws.cell(row=row, column=col, value=formula)
        cell.font = F_BODY_B if bold else F_BODY
        return cell

    HIST_ROW_KEYS = [
        (26, "revenue"), (27, "cogs"), (30, "salaries_and_benefits"),
        (31, "rent_and_overhead"), (32, "depreciation_amortization"),
        (33, "interest_expense"), (37, "taxes"),
        (44, "cash"), (45, "accounts_receivable"), (46, "inventory"),
        (47, "ppe_closing"), (51, "accounts_payable"), (52, "debt_closing"),
        (55, "equity_capital"), (56, "retained_earnings"),
        (66, "net_earnings_cf"), (67, "depreciation_cf"), (68, "change_in_nwc_cf"),
        (72, "capex"), (76, "debt_issuance_repayment"),
        (77, "equity_issuance_repayment"),
    ]
    for row, key in HIST_ROW_KEYS:
        for i, c in enumerate(HIST_COLS):
            put_input(row, c, hist[key][i], bold=(row == 26))
    put_input(81, FIRST, hist["opening_cash_first_year"])
    put_input(96, FIRST, hist["ppe_opening_first_year"])
    put_input(102, FIRST, hist["debt_opening_first_year"])
    for c in ALL_COLS:
        put_input(21, c, DIP)

    FC_ROW_KEYS = [
        (8, "revenue_growth_pct"), (9, "cogs_pct_revenue"),
        (10, "salaries_pct_revenue"), (11, "rent_and_overhead"),
        (12, "da_pct_opening_ppe"), (13, "interest_pct_avg_debt"),
        (14, "tax_pct_ebt"), (15, "ar_days"), (16, "inventory_days"),
        (17, "ap_days"), (18, "capex"), (19, "debt_issuance_repayment"),
        (20, "equity_issuance_repayment"),
    ]
    for row, key in FC_ROW_KEYS:
        for i, c in enumerate(FC_COLS):
            put_input(row, c, fa[key][i])

    def hist_backcalc(row, col, expr, denominator):
        """Historical driver back-calcs show #DIV/0! for zero-denominator
        inputs (debt-free, no-COGS, zero-EBT years); wrap only those cells
        in IFERROR(...,"") so the audit rows stay clean. The validator
        accepts both forms."""
        f = f"={expr}" if denominator else f'=IFERROR({expr},"")'
        put_formula(row, col, f)

    for i, c in enumerate(HIST_COLS):
        cl, pl = L(c), L(c - 1)
        ebt = (hist["revenue"][i] - hist["cogs"][i]
               - hist["salaries_and_benefits"][i] - hist["rent_and_overhead"][i]
               - hist["depreciation_amortization"][i] - hist["interest_expense"][i])
        if c != FIRST:
            hist_backcalc(8, c, f"{cl}26/{pl}26-1", hist["revenue"][i - 1])
        hist_backcalc(9, c, f"{cl}27/{cl}26", hist["revenue"][i])
        hist_backcalc(10, c, f"{cl}30/{cl}26", hist["revenue"][i])
        put_formula(11, c, f"={cl}31")
        hist_backcalc(12, c, f"{cl}32/{cl}47", hist["ppe_closing"][i])
        hist_backcalc(13, c, f"{cl}33/{cl}52", hist["debt_closing"][i])
        hist_backcalc(14, c, f"{cl}37/{cl}35", ebt)
        hist_backcalc(15, c, f"{cl}45/{cl}26*365", hist["revenue"][i])
        hist_backcalc(16, c, f"{cl}46/{cl}27*365", hist["cogs"][i])
        hist_backcalc(17, c, f"{cl}51/{cl}27*365", hist["cogs"][i])
        put_formula(18, c, f"={cl}72")
        put_formula(19, c, f"={cl}76")
        put_formula(20, c, f"={cl}77")

    for c in ALL_COLS:
        cl = L(c)
        put_formula(28, c, f"={cl}26-{cl}27")
        put_formula(34, c, f"=SUM({cl}30:{cl}33)")
        put_formula(35, c, f"={cl}28-{cl}34")
        put_formula(38, c, f"={cl}35-{cl}37")
    for c in FC_COLS:
        cl, pl = L(c), L(c - 1)
        put_formula(26, c, f"={pl}26*(1+{cl}8)", bold=True)
        put_formula(27, c, f"={cl}26*{cl}9")
        put_formula(30, c, f"={cl}10*{cl}26")
        put_formula(31, c, f"={cl}11")
        put_formula(32, c, f"={cl}67")
        put_formula(33, c, f"={cl}105")
        put_formula(37, c, f"={cl}35*{cl}14")

    for c in ALL_COLS:
        cl = L(c)
        put_formula(48, c, f"=SUM({cl}44:{cl}47)")
        put_formula(53, c, f"=SUM({cl}51:{cl}52)")
        put_formula(57, c, f"=SUM({cl}55:{cl}56)")
        put_formula(58, c, f"={cl}53+{cl}57")
        put_formula(60, c, f"={cl}58-{cl}48")
    for c in FC_COLS:
        cl, pl = L(c), L(c - 1)
        put_formula(44, c, f"={cl}82")
        put_formula(45, c, f"={cl}26*{cl}15/{cl}21")
        put_formula(46, c, f"={cl}27*{cl}16/{cl}21")
        put_formula(47, c, f"={cl}99")
        put_formula(51, c, f"={cl}27*{cl}17/{cl}21")
        put_formula(52, c, f"={cl}104")
        put_formula(55, c, f"={pl}55+{cl}77")
        put_formula(56, c, f"={pl}56+{cl}38")

    for c in ALL_COLS:
        cl, pl = L(c), L(c - 1)
        put_formula(89, c, f"={cl}45")
        put_formula(90, c, f"={cl}46")
        put_formula(91, c, f"={cl}51")
        put_formula(92, c, f"={cl}89+{cl}90-{cl}91")
        put_formula(93, c, f"={cl}92-{pl}92")  # first col hits spacer C: intentional
        if c != FIRST:
            put_formula(96, c, f"={pl}99")
            put_formula(102, c, f"={pl}104")
        put_formula(99, c, f"={cl}96+{cl}97-{cl}98")
        put_formula(104, c, f"={cl}102+{cl}103")
    for c in HIST_COLS:
        cl = L(c)
        put_formula(97, c, f"={cl}72")
        put_formula(98, c, f"={cl}67")
        put_formula(103, c, f"={cl}76")
        put_formula(105, c, f"={cl}33")
    for c in FC_COLS:
        cl = L(c)
        put_formula(97, c, f"={cl}18")
        put_formula(98, c, f"={cl}96*{cl}12")
        put_formula(103, c, f"={cl}19")
        put_formula(105, c, f"=AVERAGE({cl}102,{cl}104)*{cl}13")

    for c in ALL_COLS:
        cl = L(c)
        put_formula(69, c, f"={cl}66+{cl}67-{cl}68")
        put_formula(73, c, f"=SUM({cl}72)")
        put_formula(78, c, f"=SUM({cl}76:{cl}77)")
        put_formula(80, c, f"={cl}69-{cl}73+{cl}78")
        put_formula(82, c, f"=SUM({cl}80:{cl}81)")
    for c in FC_COLS:
        cl = L(c)
        put_formula(66, c, f"={cl}38")
        put_formula(67, c, f"={cl}98")
        put_formula(68, c, f"={cl}93")
        put_formula(72, c, f"={cl}97")
        put_formula(76, c, f"={cl}103")
        put_formula(77, c, f"={cl}20")
    for c in ALL_COLS[1:]:
        put_formula(81, c, f"={L(c-1)}82")

    for c in ALL_COLS:
        cl = L(c)
        cell = put_formula(3, c, f'=IFERROR(IF(ABS({cl}60)>1,"ERROR","OK"),"ERROR")')
        cell.font = F_CHECK
        cell.alignment = Alignment(horizontal="right")

    for c in ALL_COLS:
        cl = L(c)
        put_formula(111, c, f"={cl}26")
        put_formula(112, c, f"={cl}28/{cl}26")
        put_formula(114, c, f"={cl}69")
        put_formula(115, c, f"=-{cl}73")
        put_formula(116, c, f"={cl}78")

    THOUSANDS_ROWS = [11, 15, 16, 17, 18, 19, 20, 21,
                      26, 27, 28, 30, 31, 32, 33, 34, 35, 37, 38,
                      44, 45, 46, 47, 48, 51, 52, 53, 55, 56, 57, 58,
                      66, 67, 68, 69, 72, 73, 76, 77, 78, 80, 81, 82,
                      89, 90, 91, 92, 93, 96, 97, 98, 99, 102, 103, 104, 105,
                      111, 114, 115, 116]
    for r in THOUSANDS_ROWS:
        for c in ALL_COLS:
            ws.cell(row=r, column=c).number_format = NUM_FMT
    for r in (8, 9, 10, 12, 13, 14):
        for c in ALL_COLS:
            ws.cell(row=r, column=c).number_format = PCT_FMT
    for c in ALL_COLS:
        ws.cell(row=112, column=c).number_format = GPM_FMT
        ws.cell(row=60, column=c).number_format = CHK_FMT

    for r in BOLD_TOTAL_ROWS:
        for c in ALL_COLS:
            ws.cell(row=r, column=c).font = F_BODY_B

    thin, double = Side(style="thin"), Side(style="double")
    border_map = {r: Border(top=thin) for r in (8, 28, 35, 53, 92)}
    border_map[33] = Border(bottom=thin)
    for r in (38, 48, 58):
        border_map[r] = Border(top=thin, bottom=double)
    border_map[57] = Border(top=thin, bottom=thin)
    for r, border in border_map.items():
        for c in range(1, LAST + 1):
            ws.cell(row=r, column=c).border = border

    for cols, fill, font in ((HIST_COLS, FILL_LIGHT, F_YEAR_H),
                             (FC_COLS, FILL_DARK, F_YEAR_F)):
        for c in cols:
            for r in (1, 2):
                cell = ws.cell(row=r, column=c)
                cell.fill = fill
                cell.font = font
            ws.cell(row=1, column=c).alignment = Alignment(
                horizontal="center", vertical="center")
    ws.merge_cells(start_row=1, start_column=FIRST, end_row=1, end_column=FC1 - 1)
    ws.merge_cells(start_row=1, start_column=FC1, end_row=1, end_column=LAST)

    rng = f"{L(FIRST)}3:{L(LAST)}3"
    err_rule = Rule(type="containsText", operator="containsText", text="ERROR",
                    dxf=DifferentialStyle(font=Font(color="9C0006"),
                                          fill=PatternFill(bgColor="FFC7CE")))
    err_rule.formula = [f'NOT(ISERROR(SEARCH("ERROR",{L(FIRST)}3)))']
    ok_rule = Rule(type="containsText", operator="containsText", text="OK",
                   dxf=DifferentialStyle(font=Font(color="006100")))
    ok_rule.formula = [f'NOT(ISERROR(SEARCH("OK",{L(FIRST)}3)))']
    ws.conditional_formatting.add(rng, err_rule)
    ws.conditional_formatting.add(rng, ok_rule)

    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.freeze_panes = "A4"
    for r in range(1, 141):
        ws.row_dimensions[r].height = 16
    ws.row_dimensions[2].height = 24
    for r in SECTION_ROWS:
        ws.row_dimensions[r].height = 20
    for col, w in (("A", 12.83), ("B", 12.33), ("C", 11.16), ("D", 11.66)):
        ws.column_dimensions[col].width = w

    sheet_ref = f"'{ws.title}'"
    cats = Reference(ws, min_col=FIRST, max_col=LAST, min_row=2, max_row=2)
    ch1 = BarChart()
    ch1.type = "col"
    ch1.grouping = "clustered"
    ch1.title = "Revenue & Gross Profit Margin"
    ch1.add_data(Reference(ws, min_col=FIRST, max_col=LAST, min_row=111, max_row=111),
                 titles_from_data=False, from_rows=True)
    ch1.set_categories(cats)
    ch1.series[0].tx = SeriesLabel(strRef=StrRef(f"{sheet_ref}!$A$111"))
    ln = LineChart()
    ln.add_data(Reference(ws, min_col=FIRST, max_col=LAST, min_row=112, max_row=112),
                titles_from_data=False, from_rows=True)
    ln.set_categories(cats)
    ln.series[0].tx = SeriesLabel(strRef=StrRef(f"{sheet_ref}!$A$112"))
    ln.y_axis.axId = 200
    ln.y_axis.numFmt = "0%"
    ch1.y_axis.crosses = "max"
    ch1 += ln
    ws.add_chart(ch1, "A118")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.grouping = "stacked"
    ch2.overlap = 100
    ch2.title = "Cash Flow"
    ch2.add_data(Reference(ws, min_col=FIRST, max_col=LAST, min_row=114, max_row=116),
                 titles_from_data=False, from_rows=True)
    ch2.set_categories(cats)
    for i, r in enumerate((114, 115, 116)):
        ch2.series[i].tx = SeriesLabel(strRef=StrRef(f"{sheet_ref}!$A${r}"))
    ws.add_chart(ch2, "G118")

    return {"columns": f"{L(FIRST)}-{L(LAST)}",
            "historical": f"{L(FIRST)}-{L(FC1-1)}",
            "forecast": f"{L(FC1)}-{L(LAST)}"}


def build(payload, pack, out_path):
    """Build the standalone single-sheet workbook (the core, independently
    validated deliverable). For the full best-practice workbook — cover, checks,
    dashboard, sensitivity — use scripts/build_workbook.py, which reuses
    populate_model_sheet so the model sheet is byte-for-byte the same."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Three Statement Model"
    wb.properties.title = payload.get("company", "Three Statement Model")
    info = populate_model_sheet(ws, payload, pack)
    wb.save(out_path)
    return info


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--inputs", required=True)
    ap.add_argument("--jurisdiction", help="jurisdiction pack JSON (optional)")
    ap.add_argument("--output", default="model.xlsx")
    args = ap.parse_args()

    payload = json.load(open(args.inputs))
    pack = json.load(open(args.jurisdiction)) if args.jurisdiction else None
    payload = apply_pack_defaults(payload, pack)

    errors, warnings = validate_inputs(payload)
    for w in warnings:
        print("  ⚠", w)
    if errors:
        print("INPUT VALIDATION FAILED:")
        for e in errors:
            print("  ✗", e)
        sys.exit(2)
    info = build(payload, pack, args.output)
    print(f"OK: {args.output} "
          f"(columns {info['columns']}, hist {info['historical']}, "
          f"fcst {info['forecast']}"
          + (f", jurisdiction: {pack.get('country')}" if pack else "") + ")")


if __name__ == "__main__":
    main()
